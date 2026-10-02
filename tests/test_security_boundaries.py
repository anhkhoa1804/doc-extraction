"""Small adversarial inputs exercised through the production boundaries."""
from __future__ import annotations

import io
import stat
import struct
import sys
import zipfile

import pytest

from doc_extraction import cli
from doc_extraction.config import PipelineConfig
from doc_extraction.pipelines.base import LayoutResult, OCRResult
from doc_extraction.schemas.document import Document, RunMetadata
from doc_extraction.schemas.page import Page
from doc_extraction.utils.limits import (
    ExtractionLimits,
    ResourceGuard,
    ResourceLimitExceeded,
    UnsafeDocumentError,
    snapshot_input,
)
from doc_extraction.utils.safe_io import UnsafeOutputPath, output_file, run_lock
from doc_extraction.utils.serde import read_json, write_json
from doc_extraction.utils.subprocesses import run_bounded

from .fixtures import (
    make_docx,
    make_image,
    make_pdf_with_broken_cmap,
    make_pdf_with_text,
    make_pptx,
    make_xlsx,
)


def _rejected(source, output, expected, limits=None):
    with pytest.raises(expected):
        cli.process_file(source, PipelineConfig(limits=limits or ExtractionLimits()), output_root=output)
    metadata = next(output.glob("rejected-*/metadata.json"))
    result = read_json(metadata)
    assert result["status"] == "failed"
    assert not list(output.rglob("final/document.json"))
    assert not list(output.rglob(".write-*"))
    return result


@pytest.mark.parametrize("member", [
    "../outside.xml", "/absolute.xml", "C:/outside.xml", "C:outside.xml",
    "\\\\server\\share\\file", "word\\..\\outside.xml", "word/data:stream",
])
def test_archive_paths_rejected_before_any_parser(tmp_path, monkeypatch, member):
    source = tmp_path / "unsafe.docx"
    with zipfile.ZipFile(source, "w") as z:
        z.writestr("word/document.xml", "<x/>")
        z.writestr(member, "<x/>")
    monkeypatch.setattr(cli.dispatcher, "route", lambda *_args: pytest.fail("route called"))
    result = _rejected(source, tmp_path / "output", ResourceLimitExceeded)
    assert result["resource_violation"]["limit_name"] == "archive_member_path"
    assert not (tmp_path / "outside.xml").exists()


@pytest.mark.parametrize("names", [
    ["word/document.xml", "word/document.xml"],
    ["word/name.xml", "word//name.xml"],
    ["word/name.xml", "word/./name.xml"],
    ["word/name.xml", "word\\name.xml"],
    ["word/caf\u00e9.xml", "word/cafe\u0301.xml"],
])
def test_duplicate_normalized_names_are_rejected(tmp_path, names):
    source = tmp_path / "duplicates.docx"
    with zipfile.ZipFile(source, "w") as z:
        for name in names:
            z.writestr(name, "<x/>")
    _rejected(source, tmp_path / "output", UnsafeDocumentError)


def test_archive_symlink_entry_is_rejected(tmp_path):
    source = tmp_path / "link.docx"
    info = zipfile.ZipInfo("word/link")
    info.create_system = 3
    info.external_attr = (stat.S_IFLNK | 0o777) << 16
    with zipfile.ZipFile(source, "w") as z:
        z.writestr("word/document.xml", "<x/>")
        z.writestr(info, "../../victim")
    _rejected(source, tmp_path / "output", ResourceLimitExceeded)


@pytest.mark.parametrize("limits,key,payload", [
    ({"max_archive_members": 1}, "max_archive_members", "<x/>"),
    ({"max_archive_uncompressed_bytes": 8}, "max_archive_uncompressed_bytes", "<x>long</x>"),
    ({"max_archive_compression_ratio": 2}, "max_archive_compression_ratio", "<x>" + "x" * 5000 + "</x>"),
])
def test_archive_resource_failures_at_public_boundary(tmp_path, limits, key, payload):
    source = tmp_path / "quota.docx"
    with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", payload)
        z.writestr("other.xml", "<x/>")
    result = _rejected(source, tmp_path / "output", ResourceLimitExceeded, ExtractionLimits(**limits))
    assert result["resource_violation"]["limit_name"] == key


def test_corrupt_archive_local_header_is_rejected(tmp_path):
    source = tmp_path / "corrupt.docx"
    with zipfile.ZipFile(source, "w") as z:
        z.writestr("word/document.xml", "<x/>")
    content = bytearray(source.read_bytes())
    content[:4] = b"PKxx"
    source.write_bytes(content)
    _rejected(source, tmp_path / "output", ValueError)


def test_malformed_central_directory_rejected_before_routing(tmp_path, monkeypatch):
    source = make_docx(tmp_path / "central.docx")
    content = bytearray(source.read_bytes())
    central = content.index(b"PK\x01\x02")
    content[central:central + 4] = b"PKxx"
    source.write_bytes(content)
    monkeypatch.setattr(cli.dispatcher, "route", lambda *_args: pytest.fail("invalid container reached route"))
    _rejected(source, tmp_path / "output", ValueError)


def test_encrypted_archive_flag_rejected_before_consumption(tmp_path):
    source = make_docx(tmp_path / "encrypted.docx")
    content = bytearray(source.read_bytes())
    central = content.index(b"PK\x01\x02")
    flags = struct.unpack_from("<H", content, central + 8)[0]
    struct.pack_into("<H", content, central + 8, flags | 1)
    source.write_bytes(content)
    _rejected(source, tmp_path / "output", UnsafeDocumentError)


@pytest.mark.parametrize("factory,member", [
    (make_docx, "word/document.xml"),
    (make_xlsx, "xl/workbook.xml"),
    (make_pptx, "ppt/presentation.xml"),
])
@pytest.mark.parametrize("payload", [
    b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///not-an-existing-test-secret">]><x>&e;</x>',
    b'<!DOCTYPE x [<!ENTITY a "ten"><!ENTITY b "&a;&a;&a;&a;">]><x>&b;</x>',
    '<!DOCTYPE x [<!ENTITY e SYSTEM "http://127.0.0.1:9/">]><x>&e;</x>'.encode("utf-16"),
])
def test_ooxml_dtd_rejected_at_api_boundary(tmp_path, monkeypatch, factory, member, payload):
    source = factory(tmp_path / ("xml." + {make_docx: "docx", make_xlsx: "xlsx", make_pptx: "pptx"}[factory]))
    with zipfile.ZipFile(source) as z:
        parts = [(i, z.read(i)) for i in z.infolist()]
    with zipfile.ZipFile(source, "w") as z:
        for info, data in parts:
            z.writestr(info, payload if info.filename == member else data)
    monkeypatch.setattr(cli.dispatcher, "route", lambda *_args: pytest.fail("unsafe XML reached route"))
    _rejected(source, tmp_path / "output", UnsafeDocumentError)


def test_nested_archive_is_inert_and_benign_unicode_names_allowed(tmp_path):
    source = make_docx(tmp_path / "nested.docx")
    nested = io.BytesIO()
    with zipfile.ZipFile(nested, "w") as z:
        z.writestr("../../outside", "inert")
    with zipfile.ZipFile(source, "a") as z:
        z.writestr("word/embeddings/object.zip", nested.getvalue())
        z.writestr("word/notes-caf\u00e9.xml", "<x/>")
    document = cli.process_file(source, PipelineConfig(), output_root=tmp_path / "output")
    assert document.pages
    assert not (tmp_path / "outside").exists()


@pytest.mark.parametrize("kind", ["eps", "huge-raster", "typed-xml"])
def test_ooxml_declared_parts_cannot_bypass_image_or_xml_preflight(tmp_path, monkeypatch, kind):
    import xml.etree.ElementTree as ET
    source = make_docx(tmp_path / "parts.docx")
    with zipfile.ZipFile(source) as archive:
        parts = [(i, archive.read(i)) for i in archive.infolist()]
    if kind == "typed-xml":
        payload = b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///not-a-secret">]><x>&e;</x>'
        declared = "application/xml"
    elif kind == "eps":
        payload = b"%!PS-Adobe-3.0 EPSF-3.0\n%%BoundingBox: 0 0 10 10\n"
        declared = "image/x-eps"
    else:
        payload = b"BM" + struct.pack("<IHHI", 54, 0, 0, 54) + struct.pack("<IiiHHIIiiII", 40, 10000, 10000, 1, 24, 0, 0, 0, 0, 0, 0)
        declared = "image/bmp"
    with zipfile.ZipFile(source, "w") as archive:
        for info, data in parts:
            if info.filename == "[Content_Types].xml":
                root = ET.fromstring(data)
                ET.SubElement(root, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", PartName="/word/part.bin", ContentType=declared)
                data = ET.tostring(root)
            archive.writestr(info, data)
        archive.writestr("word/part.bin", payload)
    monkeypatch.setattr(cli.dispatcher, "route", lambda *_args: pytest.fail("unsafe part reached parser"))
    _rejected(source, tmp_path / "output", ResourceLimitExceeded if kind == "huge-raster" else UnsafeDocumentError)


def test_benign_ooxml_embedded_raster_preserves_native_output(tmp_path):
    import docx
    source = make_docx(tmp_path / "picture.docx")
    picture = make_image(tmp_path / "picture.png")
    document = docx.Document(source)
    document.add_picture(str(picture))
    document.save(source)
    result = cli.process_file(source, PipelineConfig(), output_root=tmp_path / "output")
    assert result.metadata.status.value in ("success", "success_with_warnings")
    assert result.pages[0].elements[0].text == "Test Heading"
    assert result.pages[0].tables


@pytest.mark.parametrize("target", ["metadata.json", "logs/pipeline.jsonl", "assembled/page-001.json", "final/document.json", "final/document.md"])
def test_preexisting_output_symlink_cannot_overwrite_target(tmp_path, target):
    source = make_docx(tmp_path / "safe.docx")
    output = tmp_path / "output"
    victim = tmp_path / "victim"
    victim.write_text("unchanged")
    link = output / target
    link.parent.mkdir(parents=True)
    link.symlink_to(victim)
    with pytest.raises(UnsafeOutputPath):
        cli.process_file(source, PipelineConfig(), output_dir=output)
    assert victim.read_text() == "unchanged"
    assert not list(output.rglob(".write-*"))


def test_output_directory_symlink_cannot_create_escaped_files(tmp_path):
    source = make_docx(tmp_path / "safe.docx")
    external = tmp_path / "external"
    external.mkdir()
    output = tmp_path / "output"
    output.symlink_to(external, target_is_directory=True)
    with pytest.raises(UnsafeOutputPath):
        cli.process_file(source, PipelineConfig(), output_dir=output)
    assert list(external.iterdir()) == []


def test_symlink_created_between_open_and_commit_does_not_redirect_write(tmp_path):
    victim = tmp_path / "victim"
    victim.write_text("unchanged")
    destination = tmp_path / "result.json"
    with output_file(destination) as stream:
        destination.symlink_to(victim)
        stream.write("safe result")
    assert victim.read_text() == "unchanged"
    assert not destination.is_symlink()
    assert destination.read_text() == "safe result"


def test_symlink_swap_of_parent_keeps_opened_directory(tmp_path):
    parent = tmp_path / "parent"
    external = tmp_path / "external"
    external.mkdir()
    with output_file(parent / "result") as stream:
        parent.rename(tmp_path / "opened-parent")
        parent.symlink_to(external, target_is_directory=True)
        stream.write("safe result")
    assert list(external.iterdir()) == []
    assert (tmp_path / "opened-parent" / "result").read_text() == "safe result"


def test_atomic_write_failure_cleans_temporary_file(tmp_path):
    destination = tmp_path / "result"
    destination.write_text("old")
    with pytest.raises(RuntimeError), output_file(destination) as stream:
        stream.write("partial")
        raise RuntimeError("controlled failure")
    assert destination.read_text() == "old"
    assert not list(tmp_path.glob(".write-*"))


def test_failure_after_assembly_unpublishes_canonical_success(tmp_path, monkeypatch):
    source = make_docx(tmp_path / "safe.docx")
    output = tmp_path / "output"
    assemble = cli.assemble_document

    def cutoff_after_publishing(*args, **kwargs):
        assemble(*args, **kwargs)
        raise ResourceLimitExceeded("max_runtime_seconds", 1, 2, "after assembly")

    monkeypatch.setattr(cli, "assemble_document", cutoff_after_publishing)
    with pytest.raises(ResourceLimitExceeded):
        cli.process_file(source, PipelineConfig(), output_dir=output)
    assert read_json(output / "metadata.json")["status"] == "failed"
    assert not (output / "final/document.json").exists()
    assert not (output / "final/document.md").exists()


def test_concurrent_run_lock_rejects_second_writer_and_releases(tmp_path):
    source = make_docx(tmp_path / "safe.docx")
    output = tmp_path / "output"
    with run_lock(output):
        write_json(output / "metadata.json", {"sentinel": True})
        with pytest.raises(UnsafeOutputPath, match="another extraction"):
            cli.process_file(source, PipelineConfig(), output_dir=output)
        assert read_json(output / "metadata.json") == {"sentinel": True}
    assert cli.process_file(source, PipelineConfig(), output_dir=output).pages
    assert stat.S_IMODE((output / "metadata.json").stat().st_mode) == 0o600


def test_rejected_input_cannot_overwrite_active_run_metadata(tmp_path):
    source = tmp_path / "malformed.docx"
    source.write_bytes(b"invalid")
    output = tmp_path / "output"
    with run_lock(output):
        write_json(output / "metadata.json", {"sentinel": True})
        with pytest.raises(UnsafeOutputPath, match="another extraction"):
            cli.process_file(source, PipelineConfig(limits=ExtractionLimits(max_input_bytes=1)), output_dir=output)
        assert read_json(output / "metadata.json") == {"sentinel": True}


def test_input_replacement_after_validation_cannot_change_parser_input(tmp_path, monkeypatch):
    source = make_docx(tmp_path / "original.docx")
    original_route = cli.dispatcher.route

    def replace_submitted_path(snapshot, config):
        source.write_bytes(b"not a document now")
        assert snapshot != source
        return original_route(snapshot, config)

    monkeypatch.setattr(cli.dispatcher, "route", replace_submitted_path)
    document = cli.process_file(source, PipelineConfig(), output_root=tmp_path / "output")
    assert document.metadata.input_path == str(source)
    assert document.pages[0].elements[0].text == "Test Heading"


def test_private_input_snapshot_removed_on_failure(tmp_path):
    source = make_docx(tmp_path / "safe.docx")
    with pytest.raises(RuntimeError), snapshot_input(source, ResourceGuard(ExtractionLimits())) as acquired:
        assert stat.S_IMODE(acquired.stat().st_mode) == 0o600
        raise RuntimeError("controlled failure")
    assert not acquired.parent.exists()


def test_fifo_input_rejected_without_waiting_for_writer(tmp_path):
    import os
    source = tmp_path / "pipe.docx"
    os.mkfifo(source)
    _rejected(source, tmp_path / "output", UnsafeDocumentError)


def test_compare_cannot_hash_oversized_input_before_preflight(tmp_path, monkeypatch):
    source = make_docx(tmp_path / "oversized.docx")
    config = tmp_path / "config.yaml"
    config.write_text(f'output_dir: "{tmp_path / "output"}"\nlimits:\n  max_input_bytes: 1\n')
    monkeypatch.setattr(cli, "sha256_file", lambda *_args: pytest.fail("rejected input hashed"))
    assert cli.main(["compare", "--input", str(source), "--config", str(config), "--backends", "baseline"]) == 1
    metadata = next((tmp_path / "output").rglob("metadata.json"))
    assert read_json(metadata)["status"] == "failed"
    assert read_json(metadata)["resource_violation"]["limit_name"] == "max_input_bytes"


@pytest.mark.parametrize("kind", ["malformed", "truncated", "eps"])
def test_untrusted_image_rejected_without_external_command(tmp_path, monkeypatch, kind):
    source = tmp_path / "image.png"
    if kind == "truncated":
        make_image(source)
        source.write_bytes(source.read_bytes()[:45])
    else:
        source.write_bytes(b"%PS-Adobe-3.0 EPSF-3.0\n%%BoundingBox: 0 0 10 10\n" if kind == "eps" else b"invalid")
    monkeypatch.setattr(cli, "_get_component_backends", lambda *_args: pytest.fail("model invoked"))
    _rejected(source, tmp_path / "output", ValueError)


def test_large_bmp_dimensions_rejected_before_decode(tmp_path):
    source = tmp_path / "huge.bmp"
    # A tiny BMP header advertises huge raster dimensions, without allocation.
    source.write_bytes(b"BM" + struct.pack("<IHHI", 54, 0, 0, 54) + struct.pack("<IiiHHIIiiII", 40, 10000, 10000, 1, 24, 0, 0, 0, 0, 0, 0))
    _rejected(source, tmp_path / "output", ResourceLimitExceeded)


def test_pdf_fallback_cannot_swallow_resource_violation(tmp_path, monkeypatch):
    source = make_pdf_with_broken_cmap(tmp_path / "broken.pdf")

    class Backend:
        def is_available(self):
            return True
        def analyze(self, _page):
            return LayoutResult()
        def recognize(self, _page):
            return OCRResult()

    monkeypatch.setattr(cli, "_get_component_backends", lambda *_args: (Backend(), Backend(), Backend()))
    config = PipelineConfig(text_quality_max_suspicious_page_ratio=1.0, limits=ExtractionLimits(max_image_pixels=10))
    with pytest.raises(ResourceLimitExceeded):
        cli.process_file(source, config, output_root=tmp_path / "output")
    metadata = read_json(next((tmp_path / "output").glob("*/metadata.json")))
    assert metadata["status"] == "failed"
    assert metadata["resource_violation"]["limit_name"] == "max_image_pixels"


def test_subprocess_arguments_are_literal_and_output_is_bounded(tmp_path):
    marker = tmp_path / "must-not-exist"
    argument = f"$(touch {marker}); echo BAD"
    result = run_bounded([sys.executable, "-c", "import sys;print(sys.argv[1])", argument], timeout=5, max_output_bytes=4096)
    assert result.stdout.strip() == argument
    assert not marker.exists()
    with pytest.raises(ResourceLimitExceeded, match="max_ocr_output_bytes"):
        run_bounded([sys.executable, "-c", "import os;os.write(1,b'x'*10000)"], timeout=5, max_output_bytes=100)


def test_subprocess_timeout_kills_owned_descendant(tmp_path):
    marker = tmp_path / "late-child"
    child = f"import time;from pathlib import Path;time.sleep(0.6);Path({str(marker)!r}).touch()"
    parent = f"import subprocess,sys,time;subprocess.Popen([sys.executable,'-c',{child!r}]);time.sleep(10)"
    with pytest.raises(ResourceLimitExceeded, match="max_runtime_seconds"):
        run_bounded([sys.executable, "-c", parent], timeout=0.15, max_output_bytes=4096)
    import time
    time.sleep(0.7)
    assert not marker.exists()


@pytest.mark.parametrize("limit", ["max_ocr_output_bytes", "max_runtime_seconds"])
def test_ocr_hard_cutoff_is_failed_at_public_entry(tmp_path, monkeypatch, limit):
    from doc_extraction.backends import tesseract_backend
    source = make_image(tmp_path / "ocr.png")

    class Layout:
        name = "controlled-layout"
        def is_available(self):
            return True
        def analyze(self, _page):
            return LayoutResult()

    monkeypatch.setattr(tesseract_backend, "is_available", lambda: True)
    code = "import os;os.write(1,b'x'*1000)" if limit == "max_ocr_output_bytes" else "import time;time.sleep(10)"

    def controlled_command(_argv, **_kwargs):
        return run_bounded([sys.executable, "-c", code], timeout=0.1, max_output_bytes=100)

    monkeypatch.setattr(tesseract_backend, "run_bounded", controlled_command)
    monkeypatch.setattr(cli, "_get_component_backends", lambda *_args: (Layout(), tesseract_backend.TesseractBackend(), None))
    with pytest.raises(ResourceLimitExceeded):
        cli.process_file(source, PipelineConfig(), output_root=tmp_path / "output")
    metadata = read_json(next((tmp_path / "output").glob("*/metadata.json")))
    assert metadata["status"] == "failed"
    assert metadata["resource_violation"]["limit_name"] == limit


def test_cli_filename_control_characters_cannot_inject_terminal_output(tmp_path, capsys):
    source = tmp_path / "bad\n\x1b[31m.pdf"
    source.write_bytes(b"invalid")
    config = tmp_path / "config.yaml"
    config.write_text(f'output_dir: "{tmp_path / "output"}"\n')
    assert cli.main(["run", "--input", str(source), "--config", str(config)]) == 1
    captured = capsys.readouterr()
    assert "\x1b" not in captured.out + captured.err
    assert "bad\\n" in captured.out


def test_html_inspector_escapes_image_attribute(tmp_path):
    from doc_extraction.evaluation.inspect_html import render_inspection_html
    metadata = RunMetadata(input_filename="safe", input_path="safe", file_hash_sha256="", file_type="png", route="image", pipeline="baseline", backend="baseline", timestamp="test")
    document = Document(document_id="safe", metadata=metadata, pages=[Page(index=0, width=1, height=1, rendered_image_path=str(tmp_path / 'image" onerror="alert(1).png'))])
    output = render_inspection_html(document, tmp_path)
    assert 'src="image" onerror=' not in output
    assert "&quot;" in output


def test_report_escapes_source_html_and_identifier_paths(tmp_path):
    from scripts.build_failure_report import build_report
    source = make_docx(tmp_path / '<img src=x onerror=alert(1)>.docx')
    output = tmp_path / "output"
    document = cli.process_file(source, PipelineConfig(), output_root=output)
    document.document_id = "../../escape"
    document.pages[0].notes.append("SUSPECT test")
    document.pages[0].tables[0].id = "../../../../outside"
    for element in document.pages[0].elements:
        if element.table_id is not None:
            element.table_id = document.pages[0].tables[0].id
    document.pages[0].tables[0].cells[0].text = "=1+1"
    write_json(next(output.glob("*/final/document.json")), document)
    report = tmp_path / "report"
    build_report(output, report)
    html = (report / "inspection/index.html").read_text()
    assert "<img src=x" not in html
    assert "&lt;img" in html
    assert "<img src=x" not in (report / "summary.md").read_text()
    assert not (tmp_path / "outside.csv").exists()
    assert "'=1+1" in next((report / "tables").glob("*.csv")).read_text()


def test_inspect_rejects_parent_traversal(tmp_path):
    with pytest.raises(UnsafeOutputPath):
        cli.main(["inspect", "../outside", "--output-dir", str(tmp_path)])


def test_native_pdf_external_actions_and_embedded_file_are_inert(tmp_path):
    import pymupdf
    source = make_pdf_with_text(tmp_path / "actions.pdf", ["This is safe native text " * 5])
    with pymupdf.open(source) as pdf:
        pdf.embfile_add("../../escape", b"not extracted")
        pdf[0].insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(10, 10, 100, 40), "uri": "http://127.0.0.1:9/"})
        pdf.saveIncr()
    document = cli.process_file(source, PipelineConfig(), output_root=tmp_path / "output")
    assert document.pages
    assert not (tmp_path / "escape").exists()
